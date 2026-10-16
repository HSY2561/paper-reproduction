from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(r'D:\论文复现')
OUT = ROOT / '论文项目完整复盘报告_20261014.docx'
FIG = ROOT / '结果图整理_20261014'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑'); styles['Normal'].font.size = Pt(9.5)
for s in ['Title','Heading 1','Heading 2','Heading 3']:
    styles[s].font.name = 'Microsoft YaHei'; styles[s]._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑'); styles[s].font.color.rgb = RGBColor(0,0,0)
styles['Title'].font.size = Pt(22); styles['Heading 1'].font.size = Pt(15); styles['Heading 2'].font.size = Pt(12)

def shade(cell, fill='D9EAF7'):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); shade(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.bold=True
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    return t

def p(text='', bold=False):
    para=doc.add_paragraph(); run=para.add_run(text); run.bold=bold; return para

doc.add_paragraph('裂缝视觉分割与机器人修复论文项目完整复盘', style='Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
para=doc.add_paragraph('面向 Crack500、自建数据与四轮差速三轴修复平台的实验总结'); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('本报告汇总项目从论文主线设计、服务器训练、本地 RTX5060 重跑、基线选择、结构消融、输入审计、逐图诊断到当前可执行后续工作的完整证据。当前结论是：SegFormer-B0 是 Crack500 精度主线基线；已测试模块均没有通过预设验证门槛，不能声称领域 SOTA。论文应同时呈现负结果、数据流限制和机器人闭环价值。')

doc.add_heading('项目时间线', level=1)
for x in [
    '初期围绕师兄论文建立视觉分割、粗定位、精定位和修复验证的论文框架。',
    '服务器阶段完成 PaddleSeg 环境、Crack500、BiSeNetV2、SegFormer、OCRNet 等候选的训练和初步比较。',
    '迁移到本地 RTX5060 后重新固定数据清单、训练预算、冻结选模和边界 F1 计算，并修正旧公式和权重选择不一致问题。',
    'BiSeNetV2 阶段完成 C0、C1b、C2b、C4 输入尺度、多种子和 PointRend 逐图诊断。',
    'SegFormer 阶段完成 Crack500 基线、UAFM、PPContext 兼容修补和 PP-LiteSeg 对照。',
    '最新波次完成 EMAU、DualAttentionModule 和 LightHamHead 三个 18,000 iter 单变量实验。']:
    doc.add_paragraph(x, style='List Number')

doc.add_heading('一 项目目标与论文主线', level=1)
p('机器人任务是让差速四轮底盘沿长裂缝自动循迹，在合适位置停车，再用三自由度笛卡尔机构完成精确定位和修复。视觉分割是循迹、裂缝宽度测量、三轴定位和修复覆盖率评估的共同输入。论文采用“机器人需求—公开基线—单模块消融—公开模型对照—坐标标定—四轮粗定位—三轴精定位—真实修复”的主线。师兄论文的消融组织逻辑被借鉴，但没有照搬其模块和结果。')

doc.add_heading('二 数据与实验协议', level=1)
table(['项目','统一设置'],[
    ('设备','RTX5060，Windows，PaddlePaddle 3.4.0，FP32'),
    ('Crack500','train 1896，val 348，test 1124'),
    ('训练','seed42，batch4，AdamW，lr 6e-5，PolynomialDecay end_lr=0，18000 iter'),
    ('增强','缩放 0.75–1.25，400×400 crop，水平翻转，颜色扰动，Normalize'),
    ('选模','仅用 val mIoU，冻结后才能访问 test'),
    ('指标','mIoU、裂缝 IoU、Precision、Recall、F1、2px 边界 F1'),
    ('数据完整性','4488 个 train/val 文件前后哈希一致')
])
p('自建数据正式训练采用 source-frame-disjoint 清单，去除 18 个图像副本或 D4 等价副本后为 train 1400、val 177。自建 test 尚未用于当前结构筛选。服务器 RTX4090 结果仅作模块参考，不能和 RTX5060 本地论文数字混用。')

doc.add_heading('三 主要实验结果', level=1)
doc.add_heading('3.1 Crack500 公开模型 test', level=2)
table(['模型','mIoU','裂缝 IoU','Precision','Recall','F1'],[
    ('SegFormer-B0','0.769475','0.570702','0.728905','0.724478','0.726684'),
    ('OCRNet-HRNet-W18','0.771509','0.575710','0.703819','0.759782','0.730731'),
    ('BiSeNetV2-C4','0.762226','0.556612','0.737494','0.694136','0.715158')])
p('SegFormer-B0 与 OCRNet 差距很小。BiSeNetV2-C4 的输入尺度调整改善了轻量基线，但仍低于两种精度模型。PP-LiteSeg、UAFM、PPContext、C1b、C2b 和波次五候选没有通过验证筛选，因此没有反向使用 test 挑选它们。')

doc.add_heading('3.2 BiSeNetV2 C4 输入与 PointRend', level=2)
table(['实验','验证 mIoU','裂缝 IoU','F1','边界 F1','判定'],[
    ('C0 BiSeNetV2','0.810794','0.646406','0.785233','0.395146','控制组'),
    ('C1b OCR 融合','0.812851','0.650291','0.788092','0.388326','裂缝 IoU 增益不足'),
    ('C2b PointRend','0.810255','0.645492','0.784558','0.411927','边界改善但主体退化'),
    ('C4 输入/裁剪','0.815506','0.654863','0.791440','0.416531','单次 seed42 通过')])
p('C4 seed43 验证 mIoU 0.816008，seed44 为 0.812876，说明输入尺度有帮助但稳定性不足。C2b 恢复方式对照表明 nearest 插值会进一步降低区域指标；PointRend 主要修正局部边界，不能保证主体区域完整。')

doc.add_heading('3.3 SegFormer-B0 波次五单模块', level=2)
table(['模型','mIoU','裂缝 IoU','Precision','Recall','F1','边界 F1','参数量','FLOPs','结论'],[
    ('SegFormer-B0','0.828021','0.678174','0.828269','0.789134','0.808228','0.417708','基线','基线','控制组'),
    ('+ EMAU','0.826912','0.676847','0.800025','0.814678','0.807285','0.407198','3.85M','4.16G*','未通过'),
    ('+ DualAttention','0.828128','0.678855','0.809990','0.807437','0.808712','0.415447','3.98M','4.18G','门槛未通过'),
    ('+ LightHamHead','0.823293','0.669802','0.808070','0.796520','0.802253','0.397322','3.43M','1.62G*','未通过')])
p('预设门槛为裂缝 IoU 至少增加 0.005、F1 至少增加 0.003、Recall 下降不超过 0.005，且 mIoU 和边界 F1 不下降。EMAU、DualAttentionModule 和 LightHamHead 均未通过，所以没有启动联合模型。带 * 的 FLOPs 对 EMAU/NMF2D 存在未计数项，只能作为相对参考。')

doc.add_heading('四 自建数据结果', level=1)
p('自建 SegFormer S0 已完成训练和验证。SegFormer-B0+UAFM 验证 mIoU 0.877269、裂缝 IoU 0.763581、Precision 0.845317、Recall 0.887602、F1 0.865944、边界 F1 0.553868。由于 Recall 相对门槛和统一筛选条件未满足，尚未冻结为最终自建模型，也尚未打开自建 test。')

doc.add_heading('五 已完成的审计与证据', level=1)
for x in [
    '服务器 SCSegamba 发现标签值与官方阈值不匹配，旧运行标记为无效。',
    '本地 train/val 哈希核对 4488 个文件，原始数据没有被修改。',
    'PointRend 完成四模型逐图诊断、恢复方式对照和典型样本 contact sheet。',
    'C4 完成 seed42/43/44 复核，发现轻量输入改进存在种子波动。',
    '波次五完成 EMAU、DualAttention、LightHamHead 各 18000 iter 和独立 val.py 复核。',
    'PIDNet-S 与 DeepLabV3P 在当前 Windows RTX5060 cuDNN 反向路径失败，未伪造完整性能结果。']:
    doc.add_paragraph(x, style='List Bullet')

doc.add_heading('六 图像与原始结果位置', level=1)
p('已整理图像目录包含 PointRend 诊断总览和三张典型样本，以及自建 SegFormer 验证掩膜示例：')
doc.add_paragraph(str(FIG), style='Intense Quote')
for img in [FIG/'01_PointRend诊断'/'diagnosis_contact_sheet_c0_c2b.png', FIG/'01_PointRend诊断'/'20160302_155857_1281_721_c0_c2b.png', FIG/'01_PointRend诊断'/'20160316_143527_1_361_c0_c2b.png']:
    if img.exists():
        para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.add_run().add_picture(str(img), width=Inches(5.9)); doc.add_paragraph(img.name).alignment=WD_ALIGN_PARAGRAPH.CENTER
p('完整逐图 CSV/JSON、训练日志和 checkpoint 保留在以下实验目录：')
for path in [ROOT/'experiments'/'retrain_wave4_20261007', ROOT/'experiments'/'retrain_wave5_20261012', ROOT/'experiments'/'local_structure_wave2_20260930', ROOT/'experiments'/'local_structure_wave3_20261001', ROOT/'experiments'/'local_segformer_wave1_20261006']:
    doc.add_paragraph(str(path), style='List Bullet')

doc.add_heading('七 后续可执行工作', level=1)
table(['顺序','工作','目的','完成标志'],[
    ('1','输入与标签流诊断','确认细裂缝是在缩放、裁剪、特征提取还是恢复阶段丢失','输出裂缝保留率、尺寸分组、插值和 Padding 审计'),
    ('2','逐图场景统计','区分极细、宽裂缝、阴影、落叶、纹理密集和低对比度错误','形成 TP/FP/FN、边界 F1 和宽度分组表'),
    ('3','单变量输入/裁剪实验','验证高分辨率或细裂缝保留策略','只改变一个数据流因素，重新用 val 门槛筛选'),
    ('4','统一速度与参数测量','服务四轮底盘和三轴平台部署选择','同一输入、同一 GPU、同一 warm-up 记录延迟/FPS'),
    ('5','source-frame-disjoint 复核','排除源帧泄漏，确认泛化','冻结模型在互斥源帧 test 上评估'),
    ('6','机器人闭环','验证视觉指标能否转化为循迹和修复质量','标定、循迹、停靠、三轴定位、宽度和覆盖率完成')])
p('推荐论文写作顺序是先写数据与协议，再写 SegFormer 基线和负结果消融，然后写输入/标签诊断，最后写相机标定、四轮粗定位、三轴精定位和真实修复。这样可以把“没有模块稳定超过基线”的事实转化为对数据流、任务目标和部署约束的清楚讨论。')

doc.add_heading('八 当前结论', level=1)
p('当前最稳妥的论文结论是：SegFormer-B0 是 Crack500 精度主线基线，BiSeNetV2-C4 是轻量部署参考；结构堆叠尚未带来稳定提升，PointRend 主要改善边界而牺牲主体区域，输入尺度调整对轻量模型有帮助但需要多种子复核。下一阶段应优先完成数据流诊断和机器人闭环，而不是继续无条件增加注意力、Transformer、Mamba 或动态卷积模块。')

doc.save(OUT)
print(OUT)
